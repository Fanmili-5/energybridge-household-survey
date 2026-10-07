"""Extract independent empirical references; replay published comparisons.

Never run the informative example IDFs as new Standard140 test results.
New models must be developed from normative Section13 specifications.
This is a reference-data and evaluator verification, not our IDF validation.
"""
import csv
import datetime as dt
import hashlib
import io
import json
import statistics
import zipfile
from pathlib import Path
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent.parent
REPO = OUT.parents[4]
SOURCE = REPO / 'artifacts/private_research/benchmark_foundation_20261004'
ARCHIVE = SOURCE / '140-2023-B-AccompanyingFiles-022726.zip'
PREFIX = 'Accompanying Files/'


def sha(data):
    return hashlib.sha256(data).hexdigest()


def json_save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n')


def measurement(z, label):
    name = PREFIX + 'Normative Materials/' + label + '-Measurements.csv'
    raw = z.read(name)
    rows = list(csv.reader(io.StringIO(raw.decode('cp1252'))))
    assert rows[2][10] == 'Qhtr' and rows[3][10] == 'Wh'
    assert rows[0][10] == 'Output' and all(x == 'Input' for x in rows[0][1:10])
    data = []
    for row in rows[4:]:
        if not row or not row[0].strip():
            continue
        time = dt.datetime.strptime(row[0], '%m/%d/%Y %H:%M')
        # The official CSV quotes some thousands separators, e.g. "1,001.48".
        # Comma is not a decimal separator in this source's numeric convention.
        values = [float(v.replace(',', '')) for v in row[1:11]]
        data.append((time, values))
    assert len({t for t, _ in data}) == len(data)
    assert all(b[0] - a[0] == dt.timedelta(hours=1) for a, b in zip(data, data[1:]))
    # Case-specific normative windows: ET110 uses18 hours (Sections13.3.2/3);
    # ET100 uses55 hours (Sections13.3.4/5). ET100B has a final extra CSV row
    # explicitly outside the validation window. Do not score initialization.
    if label.startswith('ET110'):
        start, end, hours = dt.datetime(2000,2,10,16), dt.datetime(2000,2,11,9), 18
    else:
        start, end, hours = dt.datetime(2000,9,16,8), dt.datetime(2000,9,18,14), 55
    steady = [(t,v) for t,v in data if start <= t <= end]
    assert len(steady) == hours
    return {'member': name, 'sha256': sha(raw), 'full_rows': len(data),
            'full_start': data[0][0].isoformat(), 'full_end': data[-1][0].isoformat(),
            'steady_state_start': steady[0][0].isoformat(), 'steady_state_end': steady[-1][0].isoformat(),
            'steady_state_hours': hours, 'timezone': 'GMT+1_fixed',
            'hour_timestamp_is_preceding_hour_interval_end': True,
            'Qhtr_mean_Wh_per_hour': statistics.mean(v[9] for _, v in steady)}, dict(steady)


def published_output(z, suffix, sheet, times):
    member = next(n for n in z.namelist() if n.endswith(suffix))
    raw = z.read(member)
    w = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    selected = []
    for row in w[sheet].iter_rows(values_only=True):
        if isinstance(row[0], dt.datetime) and row[0] in times:
            assert len(row) >= 11 and all(isinstance(x, (int, float)) for x in row[1:11])
            selected.append((row[0], [float(x) for x in row[1:11]]))
    w.close()
    assert len(selected) == len(times) and len({t for t, _ in selected}) == len(times)
    assert set(t for t, _ in selected) == set(times)
    # Inputs are deliberately compared separately from Qhtr. Do not use
    # measured output to fill or tune a model input.
    residuals = [max(abs(v[j] - times[t][j]) for t, v in selected) for j in range(9)]
    return {'member': member, 'sha256': sha(raw), 'sheet': sheet,
            'reported_mean_Qhtr_Wh_per_hour': statistics.mean(v[9] for _, v in selected),
            'maximum_abs_input_residuals_order_Tguards_Tcell_Qfan_HWDsafr': residuals,
            'Qhtr_hourly_Wh': [v[9] for _, v in selected],
            'Qfan_mean_Wh': statistics.mean(v[7] for _, v in selected)}


def main():
    z = zipfile.ZipFile(ARCHIVE)
    references, comparisons = [], []
    summary_member = PREFIX + 'Informative Materials/Std140-ET_Results.xlsx'
    summary_raw = z.read(summary_member)
    summary = load_workbook(io.BytesIO(summary_raw), read_only=True, data_only=True)
    published = {(r[0], r[1]): r for r in summary['Qhtr_source_data'].iter_rows(values_only=True)
                 if r[0] in {'EnergyPlus', 'DeST'} and isinstance(r[1], str)}
    for label in ['ET110A', 'ET110B', 'ET100A', 'ET100B']:
        ref, data = measurement(z, label)
        case = label + '1'; ref['case'] = case
        # Source-characterized uncertainty for these cases, NOT a universal
        # percent-error gate for Chinese dwellings or whole-home electricity.
        uncertainty_fraction = {'ET110A': .009, 'ET110B': .010,
                                'ET100A': .007, 'ET100B': .009}[label]
        ref['measured_Qhtr_2sigma_fraction'] = uncertainty_fraction
        ref['measurement_band_Wh_per_hour'] = [ref['Qhtr_mean_Wh_per_hour'] * (1-uncertainty_fraction),
                                               ref['Qhtr_mean_Wh_per_hour'] * (1+uncertainty_fraction)]
        references.append(ref)
        for engine, suffix in [
            ('EnergyPlus', 'ET100series-Output-GMT+1-EnergyPlus (110724).xlsx'),
            ('DeST', '241205-DeST-ET100series-Output-GMT+1.xlsx')]:
            output = published_output(z, suffix, case, data)
            row = published[(engine, case)]
            # Published summaries are rounded to4 decimals; a half-last-place
            # allowance plus floating epsilon is numerical, not physical acceptance.
            assert abs(output['reported_mean_Qhtr_Wh_per_hour'] - float(row[6])) <= .00005001
            difference = output['reported_mean_Qhtr_Wh_per_hour'] - ref['Qhtr_mean_Wh_per_hour']
            # The informative measurement workbook and normative CSV have
            # different serialized precision. Preserve their difference; never
            # silently replace normative data to make a summary check pass.
            informative_difference = output['reported_mean_Qhtr_Wh_per_hour'] - float(row[2])
            assert abs(informative_difference - float(row[10])) <= .00010001
            output.update(engine=engine, case=case, delta_Wh_per_hour=difference,
                          published_informative_measurement_mean_Wh_per_hour=float(row[2]),
                          normative_CSV_minus_informative_measurement_mean_Wh_per_hour=
                          ref['Qhtr_mean_Wh_per_hour'] - float(row[2]),
                          absolute_relative_difference=abs(difference) / ref['Qhtr_mean_Wh_per_hour'],
                          mean_within_measured_2sigma_band=ref['measurement_band_Wh_per_hour'][0]
                          <= output['reported_mean_Qhtr_Wh_per_hour'] <= ref['measurement_band_Wh_per_hour'][1],
                          hourly_fraction_within_average_measurement_band=sum(
                              ref['measurement_band_Wh_per_hour'][0] <= v <= ref['measurement_band_Wh_per_hour'][1]
                              for v in output['Qhtr_hourly_Wh']) / len(data),
                          is_our_generated_IDF_result=False)
            wrong_total = output['reported_mean_Qhtr_Wh_per_hour'] + output['Qfan_mean_Wh']
            assert not (ref['measurement_band_Wh_per_hour'][0] <= wrong_total <= ref['measurement_band_Wh_per_hour'][1])
            output['negative_control_heater_plus_fan_wrong_meter_rejected'] = True
            comparisons.append(output)
    summary.close()
    result = {'source_archive_path': str(ARCHIVE), 'source_archive_sha256': sha(ARCHIVE.read_bytes()),
              'published_summary_sha256': sha(summary_raw), 'reference_cases': references,
              'published_comparisons_recomputed': comparisons, 'reference_evaluator_verified': True,
              'reference_cases_recomputed': 4, 'published_engine_case_comparisons': 8,
              'new_energyplus_runs': 0, 'new_standard140_test_case_IDFs': 0,
              'local_energyplus24_1_empirical_validation_completed': False,
              'China_household_or_national_population_validated': False,
              'informative_IDF_reuse_for_new_test_results_permitted_by_readme': False,
              'normative_measurement_output_Qhtr_must_not_be_used_as_model_input': True,
              'next_test_requirement': 'build_independent_models_from_normative_Section13_then_apply_same_evaluator; add_two_auto_surface_transfer_cases_for_full_six_case_suite'}
    json_save('ETNA_EXTERNAL_ORACLE.json', result)
    print(json.dumps({'reference_cases': 4, 'published_comparisons': 8,
                      'all_published_means_in_source_measurement_band': all(x['mean_within_measured_2sigma_band'] for x in comparisons),
                      'our_IDF_empirical_validation_completed': False}, indent=2))


if __name__ == '__main__':
    main()
